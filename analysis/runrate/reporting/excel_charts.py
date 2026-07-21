"""
Excel chart generation for RunRate trend analysis.

Creates OpenPyXL charts for daily trends, MTTR/MTBF, and time bucket analysis.
"""

from typing import Tuple

try:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.styles import Font

    CHART_AVAILABLE = True
except ImportError:
    CHART_AVAILABLE = False


def _set_series_title(series, title: str) -> None:
    """
    Set series title with compatibility for different openpyxl versions.

    Args:
        series: Chart series object
        title: Title string for the series
    """
    try:
        # Modern openpyxl requires SeriesLabel object
        series.tx = SeriesLabel(v=title)
    except (TypeError, AttributeError):
        # Fallback for older openpyxl versions
        series.title = title


def check_chart_availability() -> bool:
    """
    Check if chart creation is available.

    Returns:
        True if openpyxl.chart is available, False otherwise
    """
    return CHART_AVAILABLE


def create_mttr_mtbf_chart(
    worksheet,
    data_start_row: int,
    data_end_row: int,
    placement: str = "K7",
) -> bool:
    """
    Create combined MTTR/MTBF line chart.

    Args:
        worksheet: OpenPyXL worksheet object
        data_start_row: Starting row of data (after headers)
        data_end_row: Ending row of data
        placement: Cell reference for chart placement (default "K7")

    Returns:
        True if chart created successfully, False otherwise
    """
    if not CHART_AVAILABLE:
        print("     ⚠️ Charts not available (openpyxl.chart not installed)")
        return False

    try:
        chart = LineChart()
        chart.title = "Daily MTTR & MTBF Trends"
        chart.x_axis.title = "Date"
        chart.y_axis.title = "Time (Minutes)"
        chart.width = 15
        chart.height = 8

        # MTTR data (column 5)
        mttr_data = Reference(
            worksheet,
            min_col=5,
            min_row=data_start_row,
            max_col=5,
            max_row=data_end_row,
        )

        # MTBF data (column 6)
        mtbf_data = Reference(
            worksheet,
            min_col=6,
            min_row=data_start_row,
            max_col=6,
            max_row=data_end_row,
        )

        # Categories (column 1 - dates)
        categories = Reference(
            worksheet,
            min_col=1,
            min_row=data_start_row,
            max_col=1,
            max_row=data_end_row,
        )

        chart.add_data(mttr_data, titles_from_data=False)
        chart.add_data(mtbf_data, titles_from_data=False)
        chart.set_categories(categories)

        # Add series names (use helper for openpyxl compatibility)
        _set_series_title(chart.series[0], "MTTR")
        _set_series_title(chart.series[1], "MTBF")

        worksheet.add_chart(chart, placement)
        print("     ✅ Combined MTTR/MTBF chart created")
        return True

    except Exception as e:
        print(f"     ❌ Error creating MTTR/MTBF chart: {e}")
        return False


def create_efficiency_chart(
    worksheet,
    data_start_row: int,
    data_end_row: int,
    placement: str = "K25",
) -> bool:
    """
    Create efficiency trend line chart.

    Args:
        worksheet: OpenPyXL worksheet object
        data_start_row: Starting row of data (after headers)
        data_end_row: Ending row of data
        placement: Cell reference for chart placement (default "K25")

    Returns:
        True if chart created successfully, False otherwise
    """
    if not CHART_AVAILABLE:
        return False

    try:
        chart = LineChart()
        chart.title = "Daily Efficiency Trend"
        chart.x_axis.title = "Date"
        chart.y_axis.title = "Efficiency %"
        chart.width = 15
        chart.height = 8

        # Efficiency data (column 7)
        eff_data = Reference(
            worksheet,
            min_col=7,
            min_row=data_start_row,
            max_col=7,
            max_row=data_end_row,
        )

        # Categories (column 1 - dates)
        categories = Reference(
            worksheet,
            min_col=1,
            min_row=data_start_row,
            max_col=1,
            max_row=data_end_row,
        )

        chart.add_data(eff_data, titles_from_data=False)
        chart.set_categories(categories)

        # Add series name (use helper for openpyxl compatibility)
        _set_series_title(chart.series[0], "Efficiency %")

        worksheet.add_chart(chart, placement)
        print("     ✅ Efficiency trend chart created")
        return True

    except Exception as e:
        print(f"     ❌ Error creating efficiency chart: {e}")
        return False


def create_production_downtime_chart(
    worksheet,
    data_start_row: int,
    data_end_row: int,
    placement: str = "AA7",
) -> bool:
    """
    Create production vs downtime stacked bar chart.

    Args:
        worksheet: OpenPyXL worksheet object
        data_start_row: Starting row of data (after headers)
        data_end_row: Ending row of data
        placement: Cell reference for chart placement (default "AA7")

    Returns:
        True if chart created successfully, False otherwise
    """
    if not CHART_AVAILABLE:
        return False

    try:
        chart = BarChart()
        chart.title = "Daily Production vs Downtime"
        chart.x_axis.title = "Date"
        chart.y_axis.title = "Time (Minutes)"
        chart.width = 15
        chart.height = 8
        chart.type = "col"  # Column chart
        chart.grouping = "stacked"

        # Production time (column 8) and Downtime (column 9)
        data = Reference(
            worksheet,
            min_col=8,
            min_row=data_start_row,
            max_col=9,
            max_row=data_end_row,
        )

        # Categories (column 1 - dates)
        categories = Reference(
            worksheet,
            min_col=1,
            min_row=data_start_row,
            max_col=1,
            max_row=data_end_row,
        )

        chart.add_data(data, titles_from_data=False)
        chart.set_categories(categories)

        # Add series names (use helper for openpyxl compatibility)
        _set_series_title(chart.series[0], "Production Time")
        _set_series_title(chart.series[1], "Downtime")

        worksheet.add_chart(chart, placement)
        print("     ✅ Production vs Downtime chart created")
        return True

    except Exception as e:
        print(f"     ❌ Error creating production/downtime chart: {e}")
        return False


def create_time_bucket_chart(
    worksheet,
    main_worksheet,
    daily_data_end_row: int,
    placement: str = "AA25",
) -> bool:
    """
    Create time bucket analysis bar chart.

    Args:
        worksheet: Graphs worksheet object
        main_worksheet: Main report worksheet (contains time bucket data)
        daily_data_end_row: Last row of daily data
        placement: Cell reference for chart placement (default "AA25")

    Returns:
        True if chart created successfully, False otherwise
    """
    if not CHART_AVAILABLE or not main_worksheet:
        return False

    try:
        print("   📊 Creating time bucket distribution chart...")

        # Create time bucket data table on graphs sheet
        bucket_data_start_row = daily_data_end_row + 5

        # Add headers
        worksheet[f"A{bucket_data_start_row}"] = "Time Bucket"
        worksheet[f"A{bucket_data_start_row}"].font = Font(bold=True)
        worksheet[f"B{bucket_data_start_row}"] = "Stop Events Count"
        worksheet[f"B{bucket_data_start_row}"].font = Font(bold=True)

        # Copy time bucket data from main sheet (rows 16-25, column N)
        bucket_data_row = bucket_data_start_row + 1
        for bucket in range(1, 11):  # Buckets 1-10
            worksheet[f"A{bucket_data_row}"] = bucket
            # Get count from main sheet (N16:N25)
            main_sheet_count = main_worksheet[f"N{15 + bucket}"].value or 0
            worksheet[f"B{bucket_data_row}"] = int(main_sheet_count)
            bucket_data_row += 1

        # Create chart
        chart = BarChart()
        chart.title = "Time Bucket Analysis - Stop Events Distribution (20-min intervals before stop)"
        chart.x_axis.title = "Time Bucket (Every 20 minutes)"
        chart.y_axis.title = "Stop Events Count"
        chart.width = 15
        chart.height = 8

        # Data references
        time_bucket_data = Reference(
            worksheet,
            min_col=2,  # Column B
            min_row=bucket_data_start_row + 1,
            max_col=2,
            max_row=bucket_data_start_row + 10,
        )

        time_bucket_categories = Reference(
            worksheet,
            min_col=1,  # Column A
            min_row=bucket_data_start_row + 1,
            max_col=1,
            max_row=bucket_data_start_row + 10,
        )

        chart.add_data(time_bucket_data, titles_from_data=False)
        chart.set_categories(time_bucket_categories)

        worksheet.add_chart(chart, placement)
        print("     ✅ Time bucket chart created successfully")
        return True

    except Exception as e:
        print(f"     ❌ Error creating time bucket chart: {e}")
        return False


def create_all_trend_charts(
    graphs_worksheet,
    main_worksheet,
    data_start_row: int,
    data_end_row: int,
) -> Tuple[int, int, int, int]:
    """
    Create all trend charts on the graphs worksheet.

    Args:
        graphs_worksheet: Worksheet for charts
        main_worksheet: Main data worksheet
        data_start_row: Starting row of daily data
        data_end_row: Ending row of daily data

    Returns:
        Tuple of (mttr_mtbf_created, efficiency_created, production_created, bucket_created)
        Each value is 1 if created, 0 if failed
    """
    if not CHART_AVAILABLE:
        print("   ⚠️ Chart creation not available - openpyxl.chart not installed")
        return (0, 0, 0, 0)

    print("   📈 Creating daily trend charts...")

    results = []

    # Chart 1: MTTR/MTBF Combined
    results.append(
        1
        if create_mttr_mtbf_chart(graphs_worksheet, data_start_row, data_end_row, "K7")
        else 0
    )

    # Chart 2: Efficiency Trend
    results.append(
        1
        if create_efficiency_chart(
            graphs_worksheet, data_start_row, data_end_row, "K25"
        )
        else 0
    )

    # Chart 3: Production vs Downtime
    results.append(
        1
        if create_production_downtime_chart(
            graphs_worksheet, data_start_row, data_end_row, "AA7"
        )
        else 0
    )

    # Chart 4: Time Bucket Analysis
    results.append(
        1
        if create_time_bucket_chart(
            graphs_worksheet, main_worksheet, data_end_row, "AA25"
        )
        else 0
    )

    total_created = sum(results)
    print(f"   ✅ Created {total_created}/4 trend charts")

    return tuple(results)
